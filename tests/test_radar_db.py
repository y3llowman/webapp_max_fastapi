"""Сквозные проверки радара на настоящем PostgreSQL: обязанности → напоминания → кнопки, реестр МСП,
документы, API мини-приложения. MAX не вызывается — отправка подменена.

Нужна пустая тестовая база, её схема пересоздаётся:
  TEST_DATABASE_URL=postgresql+asyncpg://max@localhost:5544/maxtest python -m unittest discover -s tests
Без TEST_DATABASE_URL тесты пропускаются.
"""
import io
import os
import sys
import unittest
from dataclasses import asdict, replace
from datetime import date, datetime, time, timedelta
from pathlib import Path
from unittest.mock import AsyncMock, patch

TEST_DB = os.environ.get("TEST_DATABASE_URL")
if TEST_DB:
    os.environ["DATABASE_URL"] = TEST_DB  # до импорта databases: движок создаётся при импорте
os.environ.setdefault("MAX_TOKEN", "test")
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from docx import Document  # noqa: E402
from httpx import ASGITransport, AsyncClient  # noqa: E402
from sqlalchemy import select  # noqa: E402

from app.api.depends import get_current_user  # noqa: E402
from data_fetching.rmsp_client import RmspRecord  # noqa: E402
from databases import Base, SessionLocal, engine, init_db  # noqa: E402
from databases.businesses_db import Business, save_business  # noqa: E402
from databases.users_db import User  # noqa: E402
from main import app  # noqa: E402
from notifications import worker  # noqa: E402
from notifications.models import BusinessProfile, Notification, RadarEvent  # noqa: E402
from notifications.planner import MSK  # noqa: E402

TODAY = date(2026, 9, 26)
MORNING = datetime(2026, 9, 26, 9, 0, tzinfo=MSK)
INN = "7707083893"
MAX_USER = 111
RECORD = RmspRecord(
    name='ООО "СЕВЕРНЫЙ ВЕТЕР"', subject_type="UL", category=1, ogrn="1027700132195", inn=INN,
    main_activity_code="41.20", main_activity_name="Строительство жилых и нежилых зданий", region_code="16",
    is_new=False, date_registered="10.08.2016", date_excluded=None, phone=None, email=None, website=None,
    has_licenses=False, is_hitech=False, is_partnership=False, is_social=False,
)


@unittest.skipUnless(TEST_DB, "нужен TEST_DATABASE_URL")
class RadarDbTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.drop_all)
        await init_db()
        async with SessionLocal() as db:
            user = User(max_user_id=MAX_USER, notification_settings={"quiet": False})
            db.add(user)
            await db.flush()
            await save_business(db, user.id, RECORD)
        self.send = AsyncMock(return_value="mid-1")
        patches = [patch.object(worker, "send_html", self.send),
                   patch.object(worker, "bot_id", AsyncMock(return_value=999))]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)

    async def asyncTearDown(self):
        await engine.dispose()  # у каждого теста свой event loop — соединения не переиспользуем

    async def answer(self, regime: str | None, headcount: int | None):
        async with SessionLocal() as db:
            profile = await db.get(BusinessProfile, INN) or BusinessProfile(inn=INN, flags={}, bank_biks=[])
            profile.tax_regime, profile.headcount = regime, headcount
            profile.has_employees = bool(headcount)
            db.add(profile)
            await db.commit()
        return await worker.materialize_for(INN, TODAY)

    async def events(self, **where) -> list[RadarEvent]:
        async with SessionLocal() as db:
            query = select(RadarEvent).filter_by(inn=INN, **where).order_by(RadarEvent.due, RadarEvent.id)
            return list(await db.scalars(query))

    async def notifications(self) -> list[Notification]:
        async with SessionLocal() as db:
            return list(await db.scalars(select(Notification).order_by(Notification.id)))

    # ---- обязанности
    async def test_profile_answers_change_obligations(self):
        added, removed = await worker.materialize_for(INN, TODAY)
        self.assertEqual((added, removed), (["Бухгалтерская отчётность"], []))

        added, removed = await self.answer("usn_income", 36)
        self.assertIn("Сведения о выполнении квоты для инвалидов", added)
        self.assertIn("Декларация по УСН", added)
        total = len(await self.events())
        await worker.materialize_for(INN, TODAY)
        self.assertEqual(len(await self.events()), total, "повторный запуск не плодит события")

        rsv = [e for e in await self.events() if e.payload["code"] == "rsv"][0]
        self.assertEqual((rsv.due, rsv.payload["shifted"], rsv.payload["original"]), (date(2026, 10, 26), True, "2026-10-25"))

        added, removed = await self.answer("usn_income", 0)
        self.assertEqual(added, [])
        self.assertIn("Расчёт по страховым взносам (РСВ)", removed)
        self.assertFalse([e for e in await self.events() if e.payload.get("code") == "rsv"])

    # ---- напоминания, отправка, кнопки
    async def test_reminders_are_grouped_and_buttons_close_tasks(self):
        await self.answer("usn_income", 36)
        await worker.queue_reminders(MORNING)
        await worker.queue_reminders(MORNING)  # дважды за день — без дублей
        pending = await self.notifications()
        self.assertTrue(pending)
        self.assertTrue(all(n.label == "T-30" for n in pending))

        await worker.dispatch()
        self.send.assert_awaited_once()  # все сроки на 26 октября — одним сообщением
        _, text, kb = self.send.await_args.args
        self.assertIn("Через 30 дней: 26 октября 2026", text)
        self.assertIn("Срок перенесён с 25 октября 2026", text)
        self.assertIn("почему вам: у вас есть сотрудники", text)
        buttons = [b for row in kb["payload"]["buttons"] for b in row]
        self.assertIn({"type": "open_app", "text": "Открыть", "contact_id": 999,
                       "payload": f"task_{pending[0].event_id}"}, buttons)
        self.assertTrue(all(n.status == "sent" for n in await self.notifications()))

        ids = [n.event_id for n in pending]
        self.assertEqual(await worker.apply_action(MAX_USER, ids, "snooze1d"), "Напомним завтра в 9:00")
        snoozed = [n for n in await self.notifications() if n.label.startswith("snooze")]
        self.assertEqual(len(snoozed), len(ids))
        tomorrow = datetime.now(MSK).date() + timedelta(days=1)  # «Завтра» — от настоящей даты
        self.assertEqual(snoozed[0].scheduled_at, datetime.combine(tomorrow, time(9), MSK))

        self.assertEqual(await worker.apply_action(MAX_USER, ids, "done"), "Отмечено как выполненное")
        self.assertEqual(await worker.apply_action(222, ids, "done"), "Задача не найдена", "чужие задачи не трогаем")
        async with SessionLocal() as db:  # напоминание по закрытой задаче не уходит
            for n in snoozed:
                (await db.get(Notification, n.id)).scheduled_at = MORNING
            await db.commit()
        self.send.reset_mock()
        await worker.dispatch()
        self.send.assert_not_awaited()
        self.assertTrue(all(n.status == "cancelled" for n in await self.notifications() if n.label.startswith("snooze")))

    async def test_overdue_is_reminded_every_day_until_closed(self):
        await self.answer(None, None)
        buh = (await self.events())[0]
        async with SessionLocal() as db:
            (await db.get(RadarEvent, buh.id)).due = date(2026, 9, 20)
            await db.commit()
        await worker.queue_reminders(MORNING)
        await worker.queue_reminders(datetime(2026, 9, 27, 9, 0, tzinfo=MSK))
        self.assertEqual([n.label for n in await self.notifications()], ["overdue:260926", "overdue:260927"])

    async def test_chat_disabled_cancels_reminders(self):
        await self.answer("usn_income", 36)
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
            user.notification_settings = {"chat": False}
            await db.commit()
        await worker.queue_reminders(MORNING)
        await worker.dispatch()
        self.send.assert_not_awaited()

    # ---- реестр МСП
    async def test_msp_exclusion_and_category_change(self):
        await worker.registry_loaded(INN, asdict(RECORD))
        self.assertEqual(await self.events(type="msp.excluded"), [])

        excluded = asdict(replace(RECORD, date_excluded="10.07.2027"))
        self.assertEqual(len(await worker.process_msp(INN, excluded, TODAY)), 1)
        self.assertEqual(await worker.process_msp(INN, excluded, TODAY), [], "открытое состояние не дублируется")
        await worker.dispatch()
        self.assertIn("Компания исключена из реестра МСП", self.send.await_args.args[1])
        self.assertIn("10 июля 2027", self.send.await_args.args[1])

        await worker.process_msp(INN, asdict(RECORD), TODAY)
        self.assertIsNotNone((await self.events(type="msp.excluded"))[0].resolved_at)

        grown = asdict(replace(RECORD, category=2))
        [event] = await worker.process_msp(INN, grown, TODAY)
        self.assertEqual((event.type, event.payload["lost_micro"], event.due), ("msp.category_changed", True, date(2027, 1, 26)))

    async def test_demo_inspection(self):
        self.assertIn("Готово", await worker.demo_event(MAX_USER, "knm"))
        text = self.send.await_args.args[1]
        self.assertIn("Запланирована проверка", text)
        self.assertIn("Ваши права", text)

    # ---- документы и API мини-приложения
    async def test_miniapp_api_and_document(self):
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
        app.dependency_overrides[get_current_user] = lambda: user
        self.addCleanup(app.dependency_overrides.clear)
        client = AsyncClient(transport=ASGITransport(app=app), base_url="http://test")
        self.addAsyncCleanup(client.aclose)

        company = (await client.get("/api/session")).json()
        self.assertEqual((company["region"], company["category"], company["needsAnswers"]),
                         ("Республика Татарстан", "Микропредприятие", True))
        await self.answer("usn_income", 36)
        company = (await client.get("/api/company")).json()
        self.assertEqual((company["headcount"], company["needsAnswers"]), ("36–100 человек", False))

        dashboard = (await client.get("/api/dashboard")).json()
        quota = next(t for t in dashboard["tasks"] if t["title"].startswith("Сведения о выполнении квоты"))
        self.assertEqual(quota["periodicity"], "Ежемесячно")
        details = (await client.get(f"/api/tasks/{quota['id']}")).json()
        self.assertEqual([s["id"] for s in details["sections"]], ["why", "steps", "how", "law", "risks"])
        self.assertTrue(details["document"])
        self.assertEqual(details["heading"], "Сведения о выполнении квоты для инвалидов за сентябрь 2026")
        law = next(s for s in details["sections"] if s["id"] == "law")
        self.assertTrue(law["link"]["url"].startswith("https://www.consultant.ru/"))

        with patch.object(worker.bot, "send_message", AsyncMock()) as send_message:
            self.assertEqual((await client.post(f"/api/tasks/{quota['id']}/document")).status_code, 204)
        media = send_message.await_args.kwargs["attachments"][0]
        doc = Document(io.BytesIO(media.buffer))
        self.assertIn("Приказ", "\n".join(p.text for p in doc.paragraphs))
        self.assertIn("Республика Татарстан", "\n".join(p.text for p in doc.paragraphs))

        rsv = next(t for t in dashboard["tasks"] if t["title"].startswith("Расчёт по страховым"))
        self.assertEqual((await client.post(f"/api/tasks/{rsv['id']}/document")).status_code, 404)

        settings = (await client.get("/api/settings/notifications")).json()
        self.assertEqual(settings["remind"], "d30-7-1")
        settings |= {"remind": "d3-0", "quietRange": "23:00–07:00"}
        self.assertEqual((await client.put("/api/settings/notifications", json=settings)).status_code, 204)
        self.assertEqual((await client.get("/api/settings/notifications")).json()["quietRange"], "23:00–07:00")
        self.assertEqual((await client.put("/api/settings/notifications", json={"remind": "d99"})).status_code, 422)

    async def test_notice_document_has_requisites(self):
        await self.answer("usn_ie", 1)
        notice = next(e for e in await self.events() if e.payload.get("code") == "usn_notice")
        async with SessionLocal() as db:
            with patch.object(worker.bot, "send_message", AsyncMock()) as send_message:
                await worker.send_document(db, notice, MAX_USER)
        doc = Document(io.BytesIO(send_message.await_args.kwargs["attachments"][0].buffer))
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        self.assertIn(INN, cells)
        self.assertIn("18210501021011000110", cells)  # КБК УСН «доходы минус расходы»

    # ---- онбординг в боте
    async def test_bot_asks_profile_questions_then_shows_obligations(self):
        from bot import message_handler as bot_flow
        say = AsyncMock()
        with patch.object(bot_flow, "say", say), patch.object(bot_flow, "bot_id", AsyncMock(return_value=999)):
            await bot_flow.next_step(MAX_USER)
            self.assertIn("режим налогообложения", say.await_args.args[1])
            await self.answer("usn_income", None)
            await bot_flow.next_step(MAX_USER)
            self.assertIn("Сколько у вас сотрудников", say.await_args.args[1])
            await self.answer("usn_income", 1)
            await bot_flow.next_step(MAX_USER)
        summary = say.await_args.args[1]
        self.assertIn("Радар настроен: 10 обязанностей", summary)
        self.assertIn("• Декларация по УСН", summary)
        async with SessionLocal() as db:
            card = bot_flow.company_card(await db.get(Business, INN), await db.get(BusinessProfile, INN))
        self.assertIn("ОКВЭД: 41.20 — Строительство жилых и нежилых зданий", card)
        self.assertIn("Микропредприятие · 1–15 человек", card)
        self.assertIn("Регион: Республика Татарстан", card)


if __name__ == "__main__":
    unittest.main()
