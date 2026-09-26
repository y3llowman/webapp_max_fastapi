"""«Подготовить документ»: черновики .docx с реквизитами из реестра МСП и выписки ЕГРЮЛ.

Суммы бот не знает, поэтому оставляет их пустыми полями. Код документа — Obligation.document.
"""
from __future__ import annotations

import io
from datetime import date

from docx import Document
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from databases.businesses_db import Business
from notifications.models import BusinessProfile, RegistrySnapshot
from radar.deadlines import region_name
from radar.render import company_name, date_ru, fio

TITLES = {
    "notice": "Уведомление об исчисленных суммах",
    "quota_order": "Приказ о квотируемых рабочих местах",
}

KBK = {
    "ndfl": ("НДФЛ", "18210102010011000110"),
    "insurance": ("Страховые взносы", "18210201000011000160"),
    "usn_income": ("Аванс по УСН «доходы»", "18210501011011000110"),
    "usn_ie": ("Аванс по УСН «доходы минус расходы»", "18210501021011000110"),
}

BLANK = "________"


async def requisites(db: AsyncSession, inn: str) -> dict:
    business = await db.get(Business, inn)
    profile = await db.get(BusinessProfile, inn)
    snapshot = (await db.execute(
        select(RegistrySnapshot.data)
        .where(RegistrySnapshot.inn == inn, RegistrySnapshot.source == "egrul")
        .order_by(RegistrySnapshot.fetched_at.desc()).limit(1)
    )).scalar_one_or_none() or {}
    director = snapshot.get("director") or {}
    director_name = " ".join(filter(None, (director.get("surname"), director.get("name"), director.get("patronymic"))))
    return {
        "name": company_name(snapshot.get("full_name") or business.name),
        "inn": inn,
        "kpp": snapshot.get("kpp"),
        "ogrn": business.ogrn,
        "address": snapshot.get("address"),
        "director_position": (director.get("position") or "Руководитель").capitalize(),
        "director_name": fio(director_name) or None,
        "region": region_name(business.region_code),
        "tax_regime": profile.tax_regime if profile else None,
    }


def _requisites_table(doc: Document, req: dict) -> None:
    rows = [("Организация", req["name"]), ("ИНН", req["inn"]), ("КПП", req["kpp"]),
            ("ОГРН", req["ogrn"]), ("Адрес", req["address"])]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value or BLANK


def _notice(doc: Document, payload: dict, due: date, req: dict) -> None:
    doc.add_heading("Уведомление об исчисленных суммах налогов (КНД 1110355)", level=1)
    doc.add_paragraph(f"{payload['title']} · {payload['period']}. Подать до {date_ru(due)}.")
    doc.add_paragraph("Черновик: реквизиты заполнены из реестров ФНС. Впишите ОКТМО и суммы, затем "
                      "отправьте уведомление через оператора ЭДО или личный кабинет налогоплательщика.")
    _requisites_table(doc, req)
    doc.add_paragraph()
    kinds = ["ndfl", "insurance"] if payload["code"] == "ndfl_notice" else [req["tax_regime"]]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Платёж", "КБК", "ОКТМО", "Сумма, ₽")):
        cell.text = text
    for kind in kinds:
        label, kbk = KBK.get(kind, ("Налог", BLANK))
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = label, kbk, BLANK, BLANK


def _quota_order(doc: Document, payload: dict, due: date, req: dict) -> None:
    doc.add_paragraph(req["name"])
    doc.add_heading("Приказ № ____", level=1)
    doc.add_paragraph("«___» __________ 20__ г.")
    doc.add_paragraph("О выделении рабочих мест для трудоустройства инвалидов в счёт квоты").runs[0].bold = True
    region = f"закона субъекта РФ ({req['region']})" if req["region"] else "закона субъекта РФ"
    doc.add_paragraph(f"В соответствии со статьёй 38 Федерального закона от 12.12.2023 № 565-ФЗ «О занятости "
                      f"населения в Российской Федерации» и {region} о квоте для приёма на работу инвалидов")
    doc.add_paragraph("ПРИКАЗЫВАЮ:")
    for item in (
        f"Выделить в счёт квоты {BLANK} рабочих мест для трудоустройства инвалидов: "
        "размер квоты — процент от среднесписочной численности, установленный законом региона.",
        "Утвердить перечень квотируемых рабочих мест (приложение к приказу).",
        "Ежемесячно, не позднее 10-го числа, представлять сведения о выполнении квоты по форме № 7 "
        f"на портале «Работа России». Ответственный: {BLANK}.",
        "Контроль за исполнением приказа оставляю за собой.",
    ):
        doc.add_paragraph(item, style="List Number")
    doc.add_paragraph()
    doc.add_paragraph(f"{req['director_position']} ____________ {req['director_name'] or BLANK}")


def build(payload: dict, due: date, req: dict) -> tuple[str, bytes]:
    """Документ payload["document"] к сроку due: (имя файла, содержимое .docx)."""
    code = payload["document"]
    doc = Document()
    if code == "notice":
        _notice(doc, payload, due, req)
    else:
        _quota_order(doc, payload, due, req)
    buffer = io.BytesIO()
    doc.save(buffer)
    return f"{code}_{req['inn']}_{due:%Y%m%d}.docx", buffer.getvalue()
